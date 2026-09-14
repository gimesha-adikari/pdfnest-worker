from __future__ import annotations

import shutil
import subprocess
import tempfile
from pathlib import Path

import pytest
from docx import Document
from openpyxl import Workbook
from pptx import Presentation

from app.api.tools.office_to_pdf.service import OfficeToPdfService


_REAL_OFFICE_RUNTIME: bool | None = None


def _build_fixture(path, format: str) -> None:
    if format == "docx":
        document = Document()
        document.add_heading("Platen Office conversion", level=1)
        document.add_paragraph("Deterministic DOCX acceptance fixture.")
        document.save(path)
        return
    if format == "xlsx":
        workbook = Workbook()
        sheet = workbook.active
        sheet["A1"] = "Platen Office conversion"
        sheet["A2"] = "Deterministic XLSX acceptance fixture."
        workbook.save(path)
        return
    if format == "pptx":
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = "Platen Office conversion"
        slide.placeholders[1].text = "Deterministic PPTX acceptance fixture."
        presentation.save(path)
        return
    raise AssertionError(f"unsupported fixture format: {format}")


def _ensure_runtime_or_skip(source, format: str) -> None:
    global _REAL_OFFICE_RUNTIME
    if _REAL_OFFICE_RUNTIME is True:
        return
    if _REAL_OFFICE_RUNTIME is False:
        pytest.skip("LibreOffice cannot execute conversions in this local runtime")

    executable = shutil.which("soffice") or shutil.which("libreoffice")
    if executable is None:
        pytest.fail("LibreOffice is required for the deterministic Office-to-PDF acceptance test")

    with tempfile.TemporaryDirectory(prefix="office-runtime-probe-") as probe_dir:
        profile = f"{probe_dir}/profile"
        output = subprocess.run(
            [
                executable,
                f"-env:UserInstallation=file://{profile}",
                "--headless",
                "--convert-to",
                {
                    "docx": "pdf:writer_pdf_Export",
                    "xlsx": "pdf:calc_pdf_Export",
                    "pptx": "pdf:impress_pdf_Export",
                }[format],
                "--outdir",
                probe_dir,
                str(source),
            ],
            capture_output=True,
            text=True,
            timeout=20,
        )
        generated = Path(probe_dir) / f"{source.stem}.pdf"
        if output.returncode != 0 or not generated.is_file() or generated.stat().st_size == 0:
            _REAL_OFFICE_RUNTIME = False
            pytest.skip(
                "LibreOffice is installed but cannot execute in this sandbox: "
                + (output.stderr or output.stdout).strip()
            )
    _REAL_OFFICE_RUNTIME = True


@pytest.mark.parametrize("format", ["docx", "xlsx", "pptx"])
def test_real_office_to_pdf_conversion(format: str, tmp_path) -> None:
    source = tmp_path / f"fixture.{format}"
    _build_fixture(source, format)
    _ensure_runtime_or_skip(source, format)
    response = OfficeToPdfService._convert_sync(format, type("Upload", (), {
        "filename": source.name,
        "file": source.open("rb"),
        "size": source.stat().st_size,
    })())
    try:
        output = response.path
        with open(output, "rb") as generated:
            assert generated.read(5) == b"%PDF-"
        assert output and response.media_type == "application/pdf"
        assert shutil.which("qpdf") is not None
        check = subprocess.run(["qpdf", "--check", output], capture_output=True, text=True)
        assert check.returncode == 0, check.stderr or check.stdout

        pdfinfo = subprocess.run(["pdfinfo", output], capture_output=True, text=True, check=True)
        pages = next(
            int(line.split(":", 1)[1].strip())
            for line in pdfinfo.stdout.splitlines()
            if line.startswith("Pages:")
        )
        assert pages == 1

        rendered_prefix = str(tmp_path / f"rendered-{format}")
        rendered = subprocess.run(
            ["pdftoppm", "-f", "1", "-singlefile", "-png", output, rendered_prefix],
            capture_output=True,
            text=True,
        )
        assert rendered.returncode == 0, rendered.stderr or rendered.stdout
        rendered_png = Path(f"{rendered_prefix}.png")
        assert rendered_png.is_file() and rendered_png.stat().st_size > 0
    finally:
        response.path and shutil.rmtree(response.path.rsplit("/", 1)[0], ignore_errors=True)
