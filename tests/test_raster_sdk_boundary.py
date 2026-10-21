"""Published, installed SDK security contract; no oversized native renders."""
import importlib
import importlib.metadata
import json
import sys
import os
import sysconfig
from pathlib import Path
from types import ModuleType, SimpleNamespace

import pytest


FACTORIES = [
    ("ocr_text_engine", "OCRTextEngineUnavailableError"),
    ("document_extraction_engine", "DocumentExtractionEngineUnavailableError"),
    ("searchable_pdf_engine", "SearchablePdfEngineUnavailableError"),
    ("ocr_markup_engine", "OcrMarkupEngineUnavailableError"),
    ("editor_ocr_engine", "EditorOcrEngineUnavailableError"),
    ("legacy_editor_ocr_engine", "LegacyEditorOcrEngineUnavailableError"),
    ("legacy_markup_ocr_engine", "LegacyMarkupOcrEngineUnavailableError"),
    ("studio_editor_extraction_engine", "StudioEditorExtractionEngineUnavailableError"),
    ("studio_markup_region_ocr_engine", "StudioMarkupRegionOcrEngineUnavailableError"),
    ("pdf_to_word_ocr_engine", "PdfToWordOcrEngineUnavailableError"),
    ("pdf_to_markdown_engine", "PdfToMarkdownEngineUnavailableError"),
]


@pytest.mark.parametrize("module,error", FACTORIES)
def test_unguarded_sdk08_cannot_construct_processor(module, error, monkeypatch):
    engine = importlib.import_module("app.core." + module)
    old_sdk = ModuleType("platen_document")
    calls = []
    old_sdk.DocumentProcessor = lambda *args, **kw: calls.append("constructed")
    old_sdk.EngineConfiguration = lambda **kw: SimpleNamespace(**kw)
    old_sdk.RasterDpiMetadataPolicy = SimpleNamespace(OMIT_DPI="omit_dpi")
    monkeypatch.setitem(sys.modules, "platen_document", old_sdk)
    with pytest.raises(getattr(engine, error)):
        engine._sdk_processor()
    assert calls == []


@pytest.mark.parametrize("module,error", FACTORIES)
def test_installed_sdk_constructs_every_factory(module, error):
    installed_sdk()
    engine = importlib.import_module("app.core." + module)
    processor = engine._sdk_processor()
    assert processor.config.raster_dpi in (144, 200)


# Literal vectors are independent of either estimator implementation.
VECTORS = [
    ((0, 0, 595.276, 841.89), 200, None, True),
    ((0, 0, 612, 792), 300, None, True),
    ((0, 0, 841.89, 595.276), 200, None, True),
    ((10.2, 20.1, 110.3, 220.4), 72, None, True),
    ((0, 0, 2498, 3998), 72, None, True),
    ((0, 0, 2499, 3998), 72, None, False),
    ((0, 0, 98, 98), 72, 9999, False),
    ((0, 0, 98, 98), 72, 10000, True),
    ((0, 0, 100_000, 100_000), 200, 100_000_000, False),
    ((0, 0, 32766, 1), 72, None, True),
    ((0, 0, 32767, 1), 72, None, False),
    ((0, 0, 1, 32767), 72, None, False),
    ((0, 0, 0, 10), 200, None, False),
    ((0, 0, -1, 10), 200, None, False),
    ((0, 0, float("nan"), 10), 200, None, False),
    ((0, 0, float("inf"), 10), 200, None, False),
    ((0, 0, 1e300, 10), 200, None, False),
    ((0, 0, 10, 10), 601, None, False),
    ((0, 0, 10, 10), True, None, False),
]


def installed_sdk():
    """Fail, rather than skip, on shared, editable or source SDK imports."""
    import platen_document
    assert not os.getenv("F11_SDK_CANDIDATE_ROOT"), "source candidate overrides are forbidden"
    assert sys.prefix != sys.base_prefix, "a disposable virtualenv is required"
    assert Path(os.environ["F11_INSTALLED_ENV_ROOT"]).resolve() == Path(sys.prefix).resolve()
    distribution = importlib.metadata.distribution("platen-document")
    assert distribution.version == "0.1.9"
    expected = Path(sysconfig.get_path("purelib")) / "platen_document/__init__.py"
    assert Path(platen_document.__file__).resolve() == expected.resolve()
    assert expected.resolve().is_relative_to(Path(sys.prefix).resolve())
    assert Path(distribution.locate_file("platen_document/__init__.py")).resolve() == expected.resolve()
    direct_url = json.loads(distribution.read_text("direct_url.json") or "{}")
    assert "dir_info" not in direct_url, "SDK source/editable installations are forbidden"
    return platen_document


def test_installed_sdk_exports_security_contract():
    sdk = installed_sdk()
    from platen_document.engine import geometry
    for name in ("RasterAdmissionError", "ConfigurationError", "OCRCancellationError"):
        assert issubclass(getattr(sdk, name), Exception)
    assert geometry.HARD_MAX_RASTER_PIXELS == 10_000_000
    assert geometry.MAX_RASTER_DPI == 600
    assert geometry.MAX_RASTER_AXIS == 32_768
    assert callable(geometry.raster_policy)


@pytest.mark.parametrize("value", [True, False, 0, -1, 1.5, float("nan"), float("inf"), "100"])
def test_installed_sdk_rejects_invalid_explicit_limits(value):
    sdk = installed_sdk()
    with pytest.raises(sdk.ConfigurationError):
        sdk.EngineConfiguration(max_raster_pixels=value)


@pytest.mark.parametrize("module,error", FACTORIES)
def test_every_factory_translates_sdk_configuration_failure(module, error, monkeypatch):
    installed_sdk()
    from app.core.ocr_v2.errors import ConfigurationError
    engine = importlib.import_module("app.core." + module)
    monkeypatch.setenv("OCR_V2_STRUCTURED_MAX_RASTER_PIXELS", "0")
    with pytest.raises(ConfigurationError):
        engine._sdk_processor()


@pytest.mark.parametrize("origin", ["checkout/src/platen_document/__init__.py", "old-worker/.venv/lib/python3.12/site-packages/platen_document/__init__.py"])
def test_installed_sdk_rejects_source_or_shared_origin(origin, tmp_path, monkeypatch):
    sdk = installed_sdk()
    monkeypatch.setattr(sdk, "__file__", str(tmp_path / origin))
    with pytest.raises(AssertionError):
        installed_sdk()


def test_installed_sdk_rejects_candidate_override(monkeypatch):
    installed_sdk()
    monkeypatch.setenv("F11_SDK_CANDIDATE_ROOT", "/old/source-candidate")
    with pytest.raises(AssertionError, match="source candidate"):
        installed_sdk()


def test_installed_sdk_rejects_editable_metadata(monkeypatch):
    installed_sdk()
    distribution = importlib.metadata.distribution("platen-document")
    monkeypatch.setattr(importlib.metadata, "distribution", lambda name: SimpleNamespace(
        version="0.1.9", locate_file=distribution.locate_file,
        read_text=lambda name: '{"dir_info": {"editable": true}}',
    ))
    with pytest.raises(AssertionError, match="source/editable"):
        installed_sdk()


@pytest.mark.parametrize("bounds,dpi,limit,admitted", VECTORS)
@pytest.mark.parametrize("rotation", [0, 90, 180, 270])
def test_installed_sdk_admission_parity(bounds, dpi, limit, admitted, rotation):
    installed_sdk()
    from platen_document.engine import geometry as sdk_geometry
    from app.core.ocr_v2 import geometry as worker_geometry
    x0, y0, x1, y1 = bounds
    page = SimpleNamespace(rect=SimpleNamespace(x0=x0,y0=y0,x1=x1,y1=y1), rotation=rotation)
    decisions = []
    for geometry in (worker_geometry, sdk_geometry):
        try:
            decision = geometry.estimate_raster(page, dpi=dpi, max_raster_pixels=limit)
        except ValueError:
            decisions.append(None)
        else:
            decisions.append((decision.width, decision.height, decision.limit))
    assert (decisions[0] is not None) is admitted
    assert decisions[0] == decisions[1]


@pytest.fixture
def rejected_page(tmp_path, monkeypatch):
    installed_sdk()
    from platen_document.engine import orchestration
    from platen_document.engine.adapters.tesseract import TesseractAdapter
    from PIL import Image
    page = SimpleNamespace(number=0, rect=SimpleNamespace(x0=0,y0=0,x1=100,y1=100,width=100,height=100), rotation=0)
    page.get_text = lambda kind: "" if kind == "text" else []
    page.get_images = lambda **kw: [(1,)]
    page.get_image_rects = lambda *args: []
    calls = []
    def forbidden(stage):
        def fail(*args, **kwargs):
            calls.append(stage)
            raise AssertionError(stage + " reached")
        return fail
    page.get_pixmap = forbidden("get_pixmap")
    monkeypatch.setattr(Image, "open", forbidden("PIL open"))
    monkeypatch.setattr(Image.Image, "convert", forbidden("PIL RGB"))
    monkeypatch.setattr(Image.Image, "save", forbidden("PNG encoding"))
    monkeypatch.setattr(TesseractAdapter, "probe", forbidden("language probe"))
    monkeypatch.setattr(TesseractAdapter, "recognize_page", forbidden("OCR adapter"))
    class Document:
        needs_pass = False
        page_count = 1
        def __len__(self): return 1
        def __getitem__(self, index): return page
        def __iter__(self): return iter([page])
        def __enter__(self): return self
        def __exit__(self, *args): pass
        def close(self): pass
    monkeypatch.setattr(orchestration.fitz, "open", lambda *args: Document())
    monkeypatch.setattr(orchestration.OCRRouter, "_available", lambda *args: True)
    source = tmp_path / "source.pdf"
    source.write_bytes(b"%PDF-small-fake")
    return source, page, calls


FAMILIES = [module for module, _ in FACTORIES] + ["ocr_markup_preview"]


def invoke_selector(module, source, output, monkeypatch, cancellation_check=None):
    """Keep selectors and the installed SDK real; isolate only input pre-scans."""
    engine = importlib.import_module("app.core." + ("ocr_markup_engine" if module == "ocr_markup_preview" else module))
    for name, value in vars(engine).items():
        if name.endswith("_ENGINE_ENV"):
            monkeypatch.setenv(value, "sdk")
    cancel = {"cancellation_check": cancellation_check}
    boxes = [{"page": 1, "x": 0, "y": 0, "width": 10, "height": 10}]
    if module == "ocr_text_engine":
        return engine.execute_ocr_text(source, language="eng", routing_policy="FAST", **cancel)
    if module == "document_extraction_engine":
        return engine.execute_document_extraction(source, language="eng", routing_policy="FAST", **cancel)
    if module == "searchable_pdf_engine":
        return engine.execute_searchable_pdf(source, output, language="eng", **cancel)
    if module == "ocr_markup_engine":
        return engine.execute_ocr_markup(source, output, action="highlight", query="word", mode="ocr", **cancel)
    if module == "ocr_markup_preview":
        return engine.execute_ocr_markup_preview(source, **cancel)
    if module == "editor_ocr_engine":
        return engine.execute_editor_ocr(source, languages=["eng"], **cancel)
    if module == "legacy_editor_ocr_engine":
        monkeypatch.setattr(engine, "_legacy_native_page_projections", lambda *args: [None])
        return engine.execute_legacy_editor_ocr(source, **cancel)
    if module == "legacy_markup_ocr_engine":
        return engine.execute_legacy_markup(source, output, boxes=boxes, action="highlight", mode="ocr", **cancel)
    if module == "studio_editor_extraction_engine":
        return engine.execute_studio_editor_extraction(source, languages=["eng"], **cancel)
    if module == "studio_markup_region_ocr_engine":
        return engine.execute_studio_markup_region_ocr(source, output, boxes=boxes, action="highlight", mode="ocr", **cancel)
    if module == "pdf_to_word_ocr_engine":
        # This selector supports cancellation through the worker context only.
        assert cancellation_check is None
        return engine.execute_pdf_to_word_ocr(source, language="eng", raster_dpis=(72,))
    if module == "pdf_to_markdown_engine":
        monkeypatch.setattr(engine, "_pdf_to_markdown_raster_dpis", lambda *args: (72,))
        return engine.execute_pdf_to_markdown(source, language="eng", routing_policy="FAST", **cancel)
    raise AssertionError("uncovered selector " + module)


@pytest.mark.parametrize("module", FAMILIES)
@pytest.mark.parametrize("limit", [100, None], ids=["smaller-context", "mandatory-default"])
def test_actual_selector_rejects_before_allocation(module, limit, rejected_page, tmp_path, monkeypatch):
    from app.core.ocr_v2.geometry import raster_policy
    from app.core.ocr_v2.errors import RasterAdmissionError
    source, page, calls = rejected_page
    if limit is None:
        page.rect = SimpleNamespace(x0=0, y0=0, x1=100_000, y1=100_000, width=100_000, height=100_000)
    output = tmp_path / "output.pdf"
    with raster_policy(limit), pytest.raises(RasterAdmissionError):
        invoke_selector(module, source, output, monkeypatch)
    assert calls == []
    assert not output.exists()


@pytest.mark.parametrize("module,via", [(module, via) for module in FAMILIES for via in ("context", "explicit") if module != "pdf_to_word_ocr_engine" or via == "context"])
def test_actual_selector_preserves_cancellation_and_resets_context(module, via, rejected_page, tmp_path, monkeypatch):
    from app.core.ocr_v2.geometry import raster_policy, _raster_context
    from platen_document.engine.geometry import _raster_context as sdk_context
    from app.jobs.cancellation import JobCancelledException
    source, page, calls = rejected_page
    original = JobCancelledException("cancelled")
    def check():
        raise original
    before_worker, before_sdk = _raster_context.get(), sdk_context.get()
    outer_check = lambda: None
    with raster_policy(200, outer_check):
        with pytest.raises(JobCancelledException) as caught:
            with raster_policy(100, check if via == "context" else None):
                invoke_selector(module, source, tmp_path / "out.pdf", monkeypatch,
                                check if via == "explicit" else None)
        assert caught.value is original
        assert _raster_context.get() == (200, outer_check)
        assert sdk_context.get() == before_sdk
    assert _raster_context.get() == before_worker
    assert calls == []


@pytest.mark.parametrize("rotation", [0, 90, 180, 270])
@pytest.mark.parametrize("clip,admitted", [((20.2,30.5,70.9,80.1), True), ((200,200,250,250), False), ((0,0,0,10), False)])
def test_installed_sdk_clip_parity(rotation, clip, admitted):
    installed_sdk()
    from platen_document.engine import geometry as sdk
    from app.core.ocr_v2 import geometry as worker
    import pymupdf as fitz
    page = SimpleNamespace(rect=fitz.Rect(10.2,20.1,110.3,220.4), rotation=rotation)
    results = []
    for geometry in (worker, sdk):
        try:
            value = geometry.estimate_raster(page, dpi=72, clip=fitz.Rect(*clip))
        except ValueError:
            results.append(None)
        else:
            results.append((value.width,value.height))
    assert (results[0] is not None) is admitted
    assert results[0] == results[1]


def test_word_selector_cannot_reclassify_typed_fatal_rejection(monkeypatch):
    from app.core import pdf_to_word_ocr_engine as engine
    from app.core.ocr_v2.errors import RasterAdmissionError
    def reject(*args, **kw):
        raise RasterAdmissionError("Document raster exceeds safety limits")
    monkeypatch.setattr(engine, "_sdk_processor", lambda **kw: SimpleNamespace(extract_document=reject))
    monkeypatch.setenv("PDF_TO_WORD_OCR_ENGINE", "sdk")
    with pytest.raises(RasterAdmissionError):
        engine.execute_pdf_to_word_ocr("fake.pdf", language="eng")


@pytest.mark.parametrize("error_name", ["RasterAdmissionError", "ConfigurationError"])
def test_legacy_markup_retains_typed_rejection(error_name, monkeypatch):
    from app.core import legacy_markup_ocr_engine as engine
    from app.core.ocr_v2 import errors
    error = getattr(errors, error_name)("safe rejection")
    def reject(*args, **kwargs):
        raise error
    monkeypatch.setattr(engine, "_sdk_processor", lambda: SimpleNamespace(extract_text=reject))
    monkeypatch.setattr(engine, "_sdk_profile", lambda: "profile")
    monkeypatch.setenv("LEGACY_MARKUP_OCR_ENGINE", "sdk")
    with pytest.raises(type(error)) as caught:
        engine.execute_legacy_markup("fake.pdf", "out.pdf", boxes=[], action="highlight", mode="ocr")
    assert caught.value is error


@pytest.mark.parametrize("selected", ["internal", "sdk"])
def test_searchable_validation_retains_cancellation_after_ocr(selected, monkeypatch):
    from app.core import searchable_pdf_engine as engine
    from app.core.ocr_v2.geometry import render_page_pixmap
    from app.jobs.cancellation import JobCancelledException
    cancelled = False
    calls = []
    page = SimpleNamespace(rect=SimpleNamespace(x0=0,y0=0,x1=10,y1=10), rotation=0)
    page.get_pixmap = lambda **kw: calls.append("native")
    def check():
        if cancelled:
            raise JobCancelledException("cancelled")
    def after(result):
        nonlocal cancelled
        cancelled = True
    def render(*args, **kwargs):
        render_page_pixmap(page, dpi=72)
    worker = SimpleNamespace(process_document=lambda *a, **kw: object())
    monkeypatch.setattr(engine, "_internal_worker", lambda: worker)
    monkeypatch.setattr(engine, "_internal_renderer", lambda: SimpleNamespace(render=render))
    monkeypatch.setattr(engine, "_sdk_processor", lambda: SimpleNamespace(
        extract_text=lambda *a, **kw: object(), make_searchable_pdf=render))
    monkeypatch.setattr(engine, "_sdk_profile", lambda: "profile")
    monkeypatch.setenv("SEARCHABLE_PDF_ENGINE", selected)
    with pytest.raises(JobCancelledException):
        engine.execute_searchable_pdf("fake.pdf", "out.pdf", language="eng",
                                      cancellation_check=check, after_ocr=after)
    assert calls == []


@pytest.mark.parametrize("selected", ["internal", "sdk"])
def test_actual_visual_validator_preserves_job_cancellation(selected, monkeypatch):
    from app.core.ocr_v2.geometry import raster_policy
    from app.core.ocr_v2.renderers import validation
    from app.jobs.cancellation import JobCancelledException
    from app.core.sdk_raster import guard_sdk_processor
    page = SimpleNamespace(rect=SimpleNamespace(x0=0,y0=0,x1=10,y1=10,width=10,height=10), rotation=0)
    page.get_images = lambda **kw: [(1,)]
    page.get_image_info = lambda: []
    page.get_text = lambda *args: []
    calls = []
    page.get_pixmap = lambda **kw: calls.append("native")
    class Document:
        def __len__(self): return 1
        def __iter__(self): return iter([page])
        def close(self): pass
    result = SimpleNamespace(pages=[SimpleNamespace(tokens=(), reading_order=())])
    def check():
        raise JobCancelledException("cancelled")
    if selected == "sdk":
        installed_sdk()
        from platen_document.engine.renderers import validation
    monkeypatch.setattr(validation.fitz, "open", lambda *a: Document())
    def validate():
        return validation.validate_searchable_pdf_artifact("source", "output", result)
    processor = guard_sdk_processor(lambda: SimpleNamespace(validate=validate)) if selected == "sdk" else SimpleNamespace(validate=validate)
    with raster_policy(None, check), pytest.raises(JobCancelledException):
        processor.validate()
    assert calls == []


@pytest.mark.parametrize("module", ["pdf_to_word_ocr_engine", "pdf_to_markdown_engine"])
def test_installed_page_scoped_preparer_retains_selected_dpi_and_worker_cap(module):
    installed_sdk()
    import io
    from PIL import Image
    from app.core.ocr_v2.geometry import raster_policy
    from app.core.ocr_v2.errors import RasterAdmissionError
    from app.core.sdk_raster import guard_sdk_processor
    page = SimpleNamespace(number=0, rotation=0,
                           rect=SimpleNamespace(x0=0,y0=0,x1=98,y1=98,width=98,height=98))
    stream = io.BytesIO()
    with Image.new("RGB", (98,98), "white") as image:
        image.save(stream, "PNG")
    calls = []
    def render(**kwargs):
        calls.append(kwargs["matrix"].a)
        return SimpleNamespace(width=98,height=98,tobytes=lambda *a: stream.getvalue())
    page.get_pixmap = render
    engine = importlib.import_module("app.core." + module)
    selected = engine._sdk_processor(raster_dpis=(72,))
    scoped = selected._structured_processor.ocr_worker.raster_preparer
    processor = guard_sdk_processor(lambda: SimpleNamespace(prepare=scoped.prepare))
    with raster_policy(9999), pytest.raises(RasterAdmissionError):
        processor.prepare(page)
    assert calls == []
    with raster_policy(10000):
        raster = processor.prepare(page)
    assert raster.dpi == 72
    assert calls == [1.0]
    raster.image.close()


@pytest.mark.parametrize("family", ["text", "searchable", "searchable_without_image", "markup", "word", "markdown"])
def test_small_native_installed_sdk_projection(family, tmp_path, monkeypatch):
    """Exercise real installed structured extraction and final consumer output."""
    installed_sdk()
    import pymupdf as fitz
    from docx import Document
    from app.api.tools.pdf_to_office.converters.word import _convert_structured_to_word
    from app.core.pdf_to_markdown_engine import execute_pdf_to_markdown

    source = tmp_path / "small.pdf"
    with fitz.open() as pdf:
        page = pdf.new_page(width=300, height=400)
        if family in {"text", "searchable"}:
            # Require real OCR on a small scan, including the source image
            # required by the existing searchable-PDF validator.
            import io
            from PIL import Image, ImageDraw, ImageFont
            stream = io.BytesIO()
            with Image.new("RGB", (600, 800), "white") as background:
                font = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", 24)
                draw = ImageDraw.Draw(background)
                for line in range(6):
                    draw.text((30, 70 + line * 80), "Ordinary approved integration text.", font=font, fill="black")
                background.save(stream, format="PNG")
            page.insert_image(page.rect, stream=stream.getvalue())
        else:
            page.insert_textbox(fitz.Rect(20, 20, 280, 380),
                                "Ordinary approved integration text. " * 20, fontsize=10)
        pdf.save(source)
    if family == "text":
        from app.core.ocr_text_engine import execute_ocr_text
        monkeypatch.setenv("OCR_TEXT_ENGINE", "sdk")
        result = execute_ocr_text(source, language="eng", routing_policy="FAST")
        assert result.pages[0].processing_source.value == "OCR_RECOGNITION"
        assert "Ordinary approved integration text" in " ".join(result.pages[0].text.split())
    elif family in {"searchable", "searchable_without_image"}:
        from app.core.searchable_pdf_engine import execute_searchable_pdf
        monkeypatch.setenv("SEARCHABLE_PDF_ENGINE", "sdk")
        output = tmp_path / "searchable.pdf"
        if family == "searchable_without_image":
            from app.core.ocr_v2.errors import RenderingNotEligibleError
            output.write_bytes(b"preexisting destination")
            with pytest.raises(RenderingNotEligibleError) as rejected:
                execute_searchable_pdf(source, output, language="eng")
            assert rejected.value.reason_code == "SOURCE_IMAGE_MISSING"
            assert output.read_bytes() == b"preexisting destination"
            assert {p.name for p in tmp_path.iterdir()} == {source.name, output.name}
            return
        result = execute_searchable_pdf(source, output, language="eng")
        assert result.pages[0].processing_source.value == "OCR_RECOGNITION"
        with fitz.open(output) as pdf:
            assert "Ordinary approved integration text" in " ".join(pdf[0].get_text().split())
    elif family == "markup":
        from app.core.ocr_markup_engine import execute_ocr_markup
        monkeypatch.setenv("OCR_MARKUP_ENGINE", "sdk")
        output = tmp_path / "marked.pdf"
        execute_ocr_markup(source, output, action="highlight", query="Ordinary", mode="smart")
        with fitz.open(output) as pdf:
            assert list(pdf[0].annots())
    elif family == "word":
        monkeypatch.setenv("PDF_TO_WORD_OCR_ENGINE", "sdk")
        output = tmp_path / "small.docx"
        _convert_structured_to_word(str(source), str(output), "eng", raster_dpis=(72,))
        assert "Ordinary approved integration text" in "\n".join(p.text for p in Document(output).paragraphs)
    else:
        monkeypatch.setenv("PDF_TO_MARKDOWN_ENGINE", "sdk")
        result = execute_pdf_to_markdown(source, language="eng", routing_policy="FAST")
        assert "Ordinary approved integration text" in result.markdown
        assert result.structured_result.pages
